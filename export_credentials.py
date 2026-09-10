"""
Export Regional Admin and Teacher login credentials to Excel.

Passwords are NOT stored in plain text anywhere (DB has SHA256 hashes), but the
scheme is deterministic:  Firstname@123#
(first word of the name, first letter capital + rest small, then @123#)
so the plain password can be recomputed from the name.

Usage:
    python export_credentials.py

Output:
    credentials_export.xlsx
        Sheet 1: "Regional Admins"  - Sr, Name, Mobile Number, Password
        Sheet 2: "Teachers"         - Sr, Name, Mobile Number, Password
"""
import os
import sys

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'EkSeSreshtha.settings')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
django.setup()

from APIS.models import Role, Teacher, User  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

OUTPUT_FILE = 'credentials_export.xlsx'


def compute_password(full_name: str) -> str:
    """Same scheme as import: first word, first letter capital + rest small + @123#"""
    first = (full_name or '').strip().split()[0] if (full_name or '').strip() else 'User'
    return f'{first.capitalize()}@123#'


def write_sheet(ws, rows):
    headers = ['Sr', 'Name', 'Mobile Number', 'Password']
    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(color='FFFFFF', bold=True)

    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for i, (name, phone, password) in enumerate(rows, start=1):
        ws.cell(row=i + 1, column=1, value=i).alignment = Alignment(horizontal='center')
        ws.cell(row=i + 1, column=2, value=name)
        # force text format so leading zeros of mobile numbers are kept
        c = ws.cell(row=i + 1, column=3, value=str(phone))
        c.number_format = '@'
        ws.cell(row=i + 1, column=4, value=password)

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 16
    ws.freeze_panes = 'A2'


def main():
    ra_role = Role.objects.filter(role_code='REGIONAL_ADMIN').first()

    # ---- Regional Admins: User accounts with REGIONAL_ADMIN role ----
    ra_users = (User.objects
                .filter(role=ra_role, status=True)
                .exclude(name__istartswith='DUMMY')
                .order_by('name')
                .values_list('name', 'phone_number'))

    # ---- Teachers: Teacher rows -> linked user account ----
    teacher_users = (Teacher.objects
                     .filter(status=True, user__status=True)
                     .exclude(user__name__istartswith='DUMMY')
                     .select_related('user')
                     .order_by('user__name')
                     .values_list('user__name', 'user__phone_number'))

    wb = Workbook()
    ws_ra = wb.active
    ws_ra.title = 'Regional Admins'
    write_sheet(ws_ra, [(n, p, compute_password(n)) for n, p in ra_users])

    ws_t = wb.create_sheet('Teachers')
    write_sheet(ws_t, [(n, p, compute_password(n)) for n, p in teacher_users])

    wb.save(OUTPUT_FILE)
    print(f'Generated {OUTPUT_FILE}')
    print(f'  Regional Admins : {ws_ra.max_row - 1}')
    print(f'  Teachers        : {ws_t.max_row - 1}')


if __name__ == '__main__':
    main()
